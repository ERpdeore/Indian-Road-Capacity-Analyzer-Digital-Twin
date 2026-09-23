"""
Defect coordinate Excel log
============================
Writes a human-readable .xlsx audit trail of every YOLOv8-detected
defect that feeds into the RoadRunner 3D digital twin -- one row per
detection, with the same lateral (t) and longitudinal (s) OpenDRIVE
coordinates that `roadrunner_export.py` actually places in the .xodr,
so the sheet and the 3D scene always agree with each other.

This does NOT replace roadrunner_export.py's JSON/.xodr pipeline --
that's still what MATLAB reads. This is an additional, readable record
of "what got detected, where" for your report/viva, and a place to
sanity-check coordinates before they go into the 3D scene.

Requires: openpyxl  (add to requirements.txt if not already there)
"""

from __future__ import annotations
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from .roadrunner_export import DEFAULT_SEGMENT_LENGTH_M, _flatten_defects


def _defect_coords(result: dict, segment_length_m: float = DEFAULT_SEGMENT_LENGTH_M) -> list[dict]:
    """Recomputes the same s/t placement _add_defect_objects() uses in
    roadrunner_export.py, so these rows match the .xodr exactly."""
    cfg = result.get("road_config", {}) or {}
    total_width_m = float(cfg.get("total_width_m", 7.0))

    defects = _flatten_defects(result.get("per_defect", {}))
    n = len(defects)
    if n == 0:
        return []

    spread_m = min(segment_length_m * 0.3, 2.0 * n)
    start_s  = max(segment_length_m / 2 - spread_m / 2, 1.0)
    step_s   = (spread_m / max(n - 1, 1)) if n > 1 else 0.0

    rows = []
    for i, d in enumerate(defects):
        rows.append({
            "type":       d["type"],
            "lateral_m":  round(d["lateral_m"], 3),
            "width_m":    round(d["width_m"], 3),
            "s_m":        round(start_s + i * step_s, 2),   # along-road position in the .xodr
            "t_m":        round(d["lateral_m"] - total_width_m / 2, 2),  # centerline-relative offset in the .xodr
        })
    return rows


def export_defects_excel(result: dict, out_path: str,
                          segment_length_m: float = DEFAULT_SEGMENT_LENGTH_M) -> str:
    """Writes a two-sheet workbook: 'Defects' (one row per YOLO
    detection with its RoadRunner s/t coordinates) and 'Summary' (the
    capacity numbers for this road). Returns out_path."""

    cfg = result.get("road_config", {}) or {}
    rows = _defect_coords(result, segment_length_m=segment_length_m)

    wb = Workbook()

    # ---- Sheet 1: Defects ----
    ws = wb.active
    ws.title = "Defects"
    headers = ["#", "Defect type", "Lateral position (m from left edge)",
               "Width (m)", "RoadRunner s (m, along road)", "RoadRunner t (m, from centerline)"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")

    if rows:
        for i, d in enumerate(rows, start=1):
            ws.append([i, d["type"], d["lateral_m"], d["width_m"], d["s_m"], d["t_m"]])
    else:
        ws.append(["-", "No defects detected", "-", "-", "-", "-"])

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 24

    # ---- Sheet 2: Summary ----
    ws2 = wb.create_sheet("Summary")
    summary_fields = [
        ("Image / segment",             cfg.get("image", "-")),
        ("Total width (m)",             cfg.get("total_width_m", "-")),
        ("Number of lanes",             cfg.get("num_lanes", "-")),
        ("Chainage (m)",                cfg.get("chainage_m", "-")),
        ("Original capacity (veh/hr)",  result.get("original_capacity_vehicles_hr", "-")),
        ("Reduced capacity (veh/hr)",   result.get("reduced_capacity_vehicles_hr", "-")),
        ("Capacity loss (%)",           result.get("capacity_loss_pct", "-")),
        ("Level of Service",            result.get("level_of_service", "-")),
        ("Free-flow speed (km/h)",      result.get("free_flow_speed_kmh", "-")),
        ("Congested speed (km/h)",      (result.get("traffic_regime") or {}).get("congested_speed_kmh", "-")),
        ("Defect count",                len(rows)),
    ]
    ws2.append(["Field", "Value"])
    for cell in ws2[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for field, value in summary_fields:
        ws2.append([field, value])
    ws2.column_dimensions["A"].width = 30
    ws2.column_dimensions["B"].width = 30

    wb.save(out_path)
    return out_path
